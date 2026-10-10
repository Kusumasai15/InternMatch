"""Private, bounded document analysis. No raw text is retained in results."""
import base64
import binascii
import io
import json
import os
import re
import time
import zipfile
from uuid import UUID
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, Field, ConfigDict
from .accounts import mutation_session, current_session
from .models import ProfileText
from .database import connection

router = APIRouter(prefix="/api/resume", tags=["Student resume"])
MAX_BYTES = 5 * 1024 * 1024
CHUNK_BYTES = 2 * 1024 * 1024
# Explicit vocabulary, word boundaries and aliases prevent arbitrary word mining.
VOCABULARY = {
    "Python": ["python"], "SQL": ["sql", "postgresql", "mysql"],
    "Machine Learning": ["machine learning", "scikit-learn", "sklearn"],
    "Pandas": ["pandas"], "NumPy": ["numpy"], "Data Analysis": ["data analysis", "data analytics"],
    "Statistics": ["statistics", "statistical analysis"], "Power BI": ["power bi", "powerbi"],
    "Excel": ["excel"], "Tableau": ["tableau"], "TensorFlow": ["tensorflow"],
    "PyTorch": ["pytorch"], "NLP": ["nlp", "natural language processing"],
    "Java": ["java"], "JavaScript": ["javascript", "js"],
    "TypeScript": ["typescript"], "React": ["react", "reactjs", "react.js"],
    "Node.js": ["node.js", "nodejs"], "HTML": ["html"], "CSS": ["css"],
    "C++": ["c++"], "C#": ["c#"], "Git": ["git"], "Docker": ["docker"],
    "Linux": ["linux"], "AWS": ["aws", "amazon web services"], "Azure": ["azure"],
    "FastAPI": ["fastapi"], "Django": ["django"], "Flask": ["flask"],
    "Figma": ["figma"], "AutoCAD": ["autocad"], "MATLAB": ["matlab"],
    "SolidWorks": ["solidworks"], "VLSI": ["vlsi"], "Embedded Systems": ["embedded systems"],
    "IoT": ["iot", "internet of things"], "Cybersecurity": ["cybersecurity"],
}


def detected_skills(text):
    return [skill for skill, aliases in VOCABULARY.items() if any(
        re.search(r"(?<!\w)" + re.escape(alias) + r"(?!\w)", text, re.I) for alias in aliases)]


def extract_text(data, extension):
    try:
        if extension == "pdf":
            if not data.startswith(b"%PDF-") or b"%%EOF" not in data[-2048:]:
                raise ValueError("This is not a complete PDF document.")
            import pymupdf
            with pymupdf.open(stream=data, filetype="pdf") as document:
                if document.is_encrypted or document.needs_pass:
                    raise ValueError("Password-protected PDFs are unsupported. Upload an unencrypted copy.")
                if document.is_repaired:
                    raise ValueError("This PDF is damaged. Export a fresh PDF and try again.")
                if not 1 <= document.page_count <= 50:
                    raise ValueError("Use a resume with 1–50 pages.")
                parts, size = [], 0
                for page in document:
                    part = page.get_text()
                    size += len(part)
                    if size > 200_000:
                        raise ValueError("This document contains too much text for a resume.")
                    parts.append(part)
                text = "\n".join(parts)
        else:
            if not data.startswith(b"PK\x03\x04"):
                raise ValueError("This is not a DOCX document.")
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                names = {entry.filename for entry in entries}
                if not {"[Content_Types].xml", "word/document.xml"}.issubset(names):
                    raise ValueError("Upload a genuine DOCX document.")
                if len(entries) > 500 or sum(entry.file_size for entry in entries) > 20 * 1024 * 1024:
                    raise ValueError("This DOCX expands beyond the safe document limit.")
                if any(entry.flag_bits & 1 or "vbaproject" in entry.filename.lower() for entry in entries):
                    raise ValueError("Encrypted or macro-enabled documents are unsupported.")
                for entry in entries:
                    if entry.filename.endswith(".xml") or entry.filename.endswith(".rels"):
                        content = archive.read(entry).upper()
                        if b"<!DOCTYPE" in content or b"<!ENTITY" in content:
                            raise ValueError("Unsafe document content is unsupported.")
            from docx import Document
            document = Document(io.BytesIO(data))
            text = "\n".join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
        if len(text) > 200_000:
            raise ValueError("This document contains too much text for a resume.")
        if len(re.sub(r"\s", "", text)) < 20:
            raise ValueError("No usable text found. Scanned resumes need OCR; upload a text-based PDF or DOCX.")
        return text
    except ValueError:
        raise
    except Exception:
        raise ValueError("The document could not be read. Export a fresh, unencrypted PDF or DOCX.") from None


def analyze(data, extension):
    text = extract_text(data, extension)
    skills = detected_skills(text)
    # Keep project skill evidence, never project names, employers, contact details or raw sentences.
    project_lines, in_projects = [], False
    for line in text.splitlines():
        heading = line.strip().lower().rstrip(":")
        if heading in {"projects", "academic projects", "personal projects"}:
            in_projects = True
        elif heading in {"education", "experience", "skills", "certifications", "work experience"}:
            in_projects = False
        if in_projects or re.search(r"project|predictive analytics", line, re.I):
            project_lines.append(line)
    project_skills = detected_skills("\n".join(project_lines))
    branch = None
    for pattern, value in [
        (r"AI\s*(?:&|and)\s*Data Science|Artificial Intelligence.*Data Science", "Artificial Intelligence & Data Science"),
        (r"Computer Science", "Computer Science & Engineering"),
        (r"Information Technology", "Information Technology"),
        (r"Electronics.*Communication", "Electronics & Communication"),
        (r"Electrical Engineering", "Electrical Engineering"),
        (r"Mechanical Engineering", "Mechanical Engineering"),
    ]:
        if re.search(pattern, text, re.I):
            branch = value
            break
    interests = []
    if any(skill in skills for skill in ["Machine Learning", "Data Analysis", "Pandas"]):
        interests.append("Data Science")
    if any(skill in skills for skill in ["React", "Node.js", "HTML"]):
        interests.append("Web Development")
    return {"skills": skills, "projectSkills": project_skills, "suggestedBranch": branch,
            "suggestedInterests": interests, "sectionsDetected": {section: bool(re.search(pattern, text, re.I)) for section, pattern in {
                "education": r"education|b\.?tech|bachelor", "projects": r"projects?|predictive analytics",
                "experience": r"experience|employment|internship", "certifications": r"certifications?|certified"}.items()},
            "appliedSkills": [], "analyzedAt": int(time.time())}


def student(session):
    if session["role"] != "student":
        raise HTTPException(403, "Resume analysis is available to student accounts only")
    return session["id"]


def save_result(db, user_id, result):
    db.execute("BEGIN IMMEDIATE")
    lock = " FOR UPDATE" if os.getenv("DATABASE_URL", "").strip() else ""
    row = db.execute("SELECT profile FROM users WHERE id=?" + lock, (user_id,)).fetchone()
    old = db.execute("SELECT details FROM resume_results WHERE user_id=?", (user_id,)).fetchone()
    previous = json.loads(old["details"]) if old else {}
    profile = json.loads(row["profile"]) if row and row["profile"] else None
    if profile:
        removed = {value.lower() for value in previous.get("appliedSkills", [])}
        baseline = [skill for skill in profile["skills"] if skill.lower() not in removed]
        existing = {value.lower() for value in baseline}
        introduced = [value for value in result["skills"] if value.lower() not in existing]
        profile["skills"] = baseline + introduced
        if len(profile["skills"]) > 100:
            raise HTTPException(422, "Keep your combined profile skills within 100 entries")
        result["appliedSkills"] = introduced
        db.execute("UPDATE users SET profile=? WHERE id=?", (json.dumps(profile), user_id))
    db.execute("INSERT INTO resume_results VALUES (?,?) ON CONFLICT(user_id) DO UPDATE SET details=excluded.details", (user_id, json.dumps(result)))
    return {"analysis": result, "profile": profile}


class Chunk(BaseModel):
    model_config = ConfigDict(extra="forbid")
    uploadId: UUID
    extension: str = Field(pattern=r"^(pdf|docx)$")
    index: int = Field(ge=0, le=2)
    data: str = Field(min_length=1, max_length=2_796_204)
    final: bool = False


@router.post("/upload")
def upload(payload: Chunk, session=Depends(mutation_session)):
    user_id = student(session)
    try:
        chunk = base64.b64decode(payload.data, validate=True)
    except (ValueError, binascii.Error):
        raise HTTPException(422, "Invalid upload encoding")
    if not chunk or len(chunk) > CHUNK_BYTES:
        raise HTTPException(413, "Upload chunks must be at most 2 MB")
    now = int(time.time())
    with connection() as db:
        db.execute("BEGIN IMMEDIATE")
        db.execute("DELETE FROM resume_uploads WHERE expires_at<?", (now,))
        if payload.index == 0:
            key = "resume:" + user_id
            db.execute("DELETE FROM auth_attempts WHERE key=? AND started_at<?", (key, now - 900))
            db.execute("INSERT INTO auth_attempts VALUES (?,1,?) ON CONFLICT(key) DO UPDATE SET attempts=auth_attempts.attempts+1", (key, now))
            if db.execute("SELECT attempts FROM auth_attempts WHERE key=?", (key,)).fetchone()["attempts"] > 5:
                raise HTTPException(429, "Try again in 15 minutes; at most five resume uploads are allowed per window")
        lock = " FOR UPDATE" if os.getenv("DATABASE_URL", "").strip() else ""
        row = db.execute("SELECT * FROM resume_uploads WHERE user_id=?" + lock, (user_id,)).fetchone()
        if payload.index == 0:
            pieces, total = [], 0
        elif not row or row["upload_id"] != str(payload.uploadId) or row["extension"] != payload.extension:
            raise HTTPException(409, "Upload expired or changed. Select the file again.")
        else:
            pieces, total = json.loads(row["chunks"]), row["size"]
        if len(pieces) != payload.index:
            raise HTTPException(409, "Upload chunks arrived out of order. Try again.")
        total += len(chunk)
        if total > MAX_BYTES:
            raise HTTPException(413, "Resumes must be at most 5 MB")
        pieces.append(payload.data)
        if payload.final:
            db.execute("DELETE FROM resume_uploads WHERE user_id=?", (user_id,))
        else:
            db.execute("INSERT INTO resume_uploads VALUES (?,?,?,?,?,?) ON CONFLICT(user_id) DO UPDATE SET upload_id=excluded.upload_id,extension=excluded.extension,chunks=excluded.chunks,size=excluded.size,expires_at=excluded.expires_at", (user_id, str(payload.uploadId), payload.extension, json.dumps(pieces), total, now + 600))
    if not payload.final:
        return {"received": total}
    try:
        result = analyze(b"".join(base64.b64decode(piece) for piece in pieces), payload.extension)
    except ValueError as error:
        raise HTTPException(422, str(error))
    with connection() as db:
        return save_result(db, user_id, result)


@router.get("")
def get_result(session=Depends(current_session)):
    with connection() as db:
        db.execute("DELETE FROM resume_uploads WHERE expires_at<?", (int(time.time()),))
        row = db.execute("SELECT details FROM resume_results WHERE user_id=?", (student(session),)).fetchone()
    return {"analysis": json.loads(row["details"]) if row else None}


class Review(BaseModel):
    model_config = ConfigDict(extra="forbid")
    skills: list[ProfileText] = Field(max_length=100)


@router.put("")
def review(payload: Review, session=Depends(mutation_session)):
    user_id = student(session)
    with connection() as db:
        row = db.execute("SELECT details FROM resume_results WHERE user_id=?", (user_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Upload a resume first")
        result = json.loads(row["details"])
        result["skills"] = list({value.lower(): value for value in payload.skills}.values())
        result["projectSkills"] = [value for value in result["projectSkills"] if value in result["skills"]]
        return save_result(db, user_id, result)


@router.delete("")
def delete_result(session=Depends(mutation_session)):
    user_id = student(session)
    with connection() as db:
        row = db.execute("SELECT details FROM resume_results WHERE user_id=?", (user_id,)).fetchone()
        profile = None
        if row:
            profile = save_result(db, user_id, {"skills": [], "appliedSkills": []})["profile"]
        db.execute("DELETE FROM resume_results WHERE user_id=?", (user_id,))
        db.execute("DELETE FROM resume_uploads WHERE user_id=?", (user_id,))
    return {"analysis": None, "profile": profile}
